from datetime import date

from hauslog.alt_import import AltEintrag, alt_daten_importieren
from hauslog.daten import Datenbank
from hauslog.export import csv_erzeugen, xlsx_erzeugen
from hauslog.formate import mit_einheit, zahl_lesen


def test_loeschen_eines_zaehlers_loescht_seine_ablesungen():
    db = Datenbank()
    zid = db.zaehler_speichern(name="Strom", art="strom", einheit="kWh")
    db.ablesung_speichern(zaehler_id=zid, datum=date(2026, 1, 1), stand=1)
    db.zaehler_loeschen(zid)
    assert db.alle_zaehler() == []
    assert db.alle_ablesungen() == []
    # Weich gelöscht und zum Senden vorgemerkt
    assert {z.geloescht for z in db.ungesendete_zaehler()} == {True}
    assert db.anzahl_ungesendet() == 2


def test_aenderung_benachrichtigt_beobachter():
    db = Datenbank()
    aufrufe = []
    db.beobachten(lambda: aufrufe.append(1))
    db.zaehler_speichern(name="Gas", art="gas", einheit="m³")
    assert aufrufe == [1]


def test_alt_import():
    db = Datenbank()
    anzahl = alt_daten_importieren(db, [
        AltEintrag("strom", "2024-01-15", 312.5, "HZ"),
        AltEintrag("strom", "2024-02-15", 600, ""),
        AltEintrag("gas", "2024-01-15", 123.4, ""),
        AltEintrag("gas", "kaputt", 1, ""),
    ])  # fmt: skip
    assert anzahl == 3
    zaehler = db.alle_zaehler()
    assert [(z.name, z.einheit) for z in zaehler] == [("Gas", "m³"), ("Strom", "kWh")]
    strom = next(z for z in zaehler if z.art == "strom")
    assert [a.stand for a in reversed(db.ablesungen(strom.id))] == [312.5, 600]


def _beispiel() -> Datenbank:
    db = Datenbank()
    zid = db.zaehler_speichern(name="Strom; Haus", art="strom", einheit="kWh")
    db.ablesung_speichern(zaehler_id=zid, datum=date(2026, 2, 1), stand=1124.5)
    db.ablesung_speichern(zaehler_id=zid, datum=date(2026, 1, 1), stand=1000)
    return db


def test_csv_export():
    db = _beispiel()
    daten = csv_erzeugen(db.alle_zaehler(), db.alle_ablesungen())
    assert daten.startswith(b"\xef\xbb\xbf")
    assert daten.decode("utf-8-sig").splitlines() == [
        "Zähler;Art;Zählernummer;Datum;Stand;Einheit;Verbrauch;Notiz",
        '"Strom; Haus";Strom;;2026-01-01;1000,0;kWh;;',
        '"Strom; Haus";Strom;;2026-02-01;1124,5;kWh;124,5;',
    ]


def test_xlsx_export():
    import io

    from openpyxl import load_workbook

    db = _beispiel()
    mappe = load_workbook(
        io.BytesIO(xlsx_erzeugen(db.alle_zaehler(), db.alle_ablesungen()))
    )
    assert mappe.sheetnames == ["Strom; Haus"]
    zeilen = list(mappe["Strom; Haus"].values)
    assert len(zeilen) == 3
    assert zeilen[2][1:3] == (1124.5, 124.5)


def test_zahlen_deutsch():
    assert mit_einheit(1124.5, "kWh") == "1.124,5 kWh"
    assert mit_einheit(3, "") == "3"
    assert zahl_lesen("1.234,5") == 1234.5
    assert zahl_lesen("1234.5") == 1234.5
    assert zahl_lesen("abc") is None
