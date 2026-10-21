"""Export als CSV (für deutsches Excel) und XLSX."""

import io
import re
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font

from .daten import Ablesung, Zaehler
from .verbrauch import Messpunkt, abschnitte_berechnen
from .zaehler_art import ZaehlerArt


def _mit_verbrauch(ablesungen: list[Ablesung]) -> list[tuple[Ablesung, float | None]]:
    """Ablesungen älteste zuerst, mit Verbrauch seit der vorherigen
    (``None`` bei der ersten oder nach Zählertausch)."""
    sortiert = sorted(ablesungen, key=lambda a: a.datum)
    abschnitte = abschnitte_berechnen(
        Messpunkt(date.fromisoformat(a.datum), a.stand) for a in sortiert
    )
    return [
        (a, None if i == 0 else abschnitte[i - 1].verbrauch)
        for i, a in enumerate(sortiert)
    ]


def _nach_zaehler(ablesungen: list[Ablesung]) -> dict[str, list[Ablesung]]:
    ergebnis: dict[str, list[Ablesung]] = {}
    for a in ablesungen:
        ergebnis.setdefault(a.zaehler_id, []).append(a)
    return ergebnis


def csv_erzeugen(zaehler: list[Zaehler], ablesungen: list[Ablesung]) -> bytes:
    """CSV für deutsches Excel: Semikolon, Dezimalkomma, UTF-8 mit BOM."""

    def feld(s: str) -> str:
        return f'"{s.replace(chr(34), chr(34) * 2)}"' if re.search('[;"\n\r]', s) else s

    def zahl(v: float | None) -> str:
        return "" if v is None else repr(float(v)).replace(".", ",")

    zeilen = ["Zähler;Art;Zählernummer;Datum;Stand;Einheit;Verbrauch;Notiz"]
    pro_zaehler = _nach_zaehler(ablesungen)
    for z in zaehler:
        for a, verbrauch in _mit_verbrauch(pro_zaehler.get(z.id, [])):
            zeilen.append(";".join([
                feld(z.name), feld(ZaehlerArt.aus_schluessel(z.art).label),
                feld(z.zaehlernummer), a.datum, zahl(a.stand), feld(z.einheit),
                zahl(verbrauch), feld(a.notiz),
            ]))  # fmt: skip
    return "﻿".encode() + ("\r\n".join(zeilen) + "\r\n").encode()


def xlsx_erzeugen(zaehler: list[Zaehler], ablesungen: list[Ablesung]) -> bytes:
    """Excel-Datei mit einem Blatt pro Zähler."""
    mappe = Workbook()
    standard = mappe.active
    pro_zaehler = _nach_zaehler(ablesungen)
    vergeben: set[str] = set()

    for z in zaehler:
        blatt = mappe.create_sheet(_blatt_name(z.name, vergeben))
        einheit = f" ({z.einheit})" if z.einheit else ""
        blatt.append(["Datum", f"Stand{einheit}", f"Verbrauch{einheit}", "Notiz"])
        for zelle in blatt[1]:
            zelle.font = Font(bold=True)
        for a, verbrauch in _mit_verbrauch(pro_zaehler.get(z.id, [])):
            blatt.append([date.fromisoformat(a.datum), a.stand, verbrauch, a.notiz])
            blatt.cell(blatt.max_row, 1).number_format = "DD.MM.YYYY"
        for spalte, breite in zip("ABCD", (14, 16, 16, 40), strict=True):
            blatt.column_dimensions[spalte].width = breite

    if zaehler:
        mappe.remove(standard)
    puffer = io.BytesIO()
    mappe.save(puffer)
    return puffer.getvalue()


def _blatt_name(name: str, vergeben: set[str]) -> str:
    """Excel erlaubt höchstens 31 Zeichen und keine Zeichen wie ``/`` oder ``:``."""
    basis = re.sub(r"[\[\]:*?/\\]", "-", name).strip()[:28] or "Zähler"
    kandidat, i = basis, 2
    while kandidat.lower() in vergeben:
        kandidat, i = f"{basis} {i}", i + 1
    vergeben.add(kandidat.lower())
    return kandidat
